"""
Cleans a legacy "FIFO Opening Quantity" export into the exact column shape
the app's real Opening Stock import template expects.

Checked against app/Services/Import/Templates/OpeningStockImportTemplate.php
directly, not guessed:
- Real required fields: Item Code, Warehouse, Date, Qty, Unit Cost.
  There is NO UOM, Description, or Item Batch field in this template at all
  -- Opening Stock has no batch tracking, and UOM/Description come from the
  Item master, not from this import.
- Header labels below match the template's field labels/synonyms exactly
  (HeaderDetector normalizes to lowercase+alnum, so "Item Code" and
  "Unit Cost" auto-map with no manual mapping step needed in the UI).
- One row = one whole Opening Stock document (header + single line) in this
  template (see persist()), so there's no need to split output per
  warehouse -- a single file with mixed warehouses/dates works fine. This
  differs from the original ask, which assumed doc = warehouse batch.

Usage: py prepare_opening_stock_import.py <source.xlsx> [output.xlsx]
"""

import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

REQUIRED_COLUMNS = ["Item Code", "Warehouse", "Date", "Qty", "Unit Cost"]
ITEM_CODE_RE = re.compile(r"^[A-Za-z0-9._/\- ]+$")


def find_header_row(ws) -> int:
    for row in ws.iter_rows(min_row=1, max_row=20):
        values = [c.value for c in row]
        if values[:3] == ["Date", "Item Code", "Description"]:
            return row[0].row
    raise ValueError("Could not find the 'Date | Item Code | Description | ...' header row in the first 20 rows.")


def parse_date(raw) -> str | None:
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, str):
        try:
            return datetime.strptime(raw.strip(), "%d/%m/%Y").date().isoformat()
        except ValueError:
            return None
    return None


def parse_number(raw) -> float | None:
    if isinstance(raw, (int, float)):
        return float(raw)
    if isinstance(raw, str):
        cleaned = raw.replace(",", "").strip()
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    src_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else src_path.with_name("OpeningStock_Import.xlsx")

    wb = openpyxl.load_workbook(src_path, data_only=True)
    ws = wb.active

    header_row = find_header_row(ws)
    col = {cell.value: cell.column for cell in ws[header_row]}

    cleaned_rows = []
    warnings = []
    unique_item_codes = set()

    for excel_row in ws.iter_rows(min_row=header_row + 1, values_only=False):
        item_code = excel_row[col["Item Code"] - 1].value
        description = excel_row[col["Description"] - 1].value
        location = excel_row[col["Location"] - 1].value
        date_raw = excel_row[col["Date"] - 1].value
        qty_raw = excel_row[col["Qty"] - 1].value
        price_raw = excel_row[col["Price"] - 1].value
        row_no = excel_row[0].row

        # The trailing "Grand Total" row carries its label in Description, not
        # Item Code (see the source file's own layout) -- detect it there, not
        # via a blank Item Code, or a real row with a genuinely missing Item
        # Code would get silently dropped instead of reported below.
        if description is not None and str(description).strip().lower() == "grand total":
            continue
        if all(v is None for v in (item_code, location, date_raw, qty_raw, price_raw)):
            continue

        item_code = str(item_code).strip() if item_code is not None else ""
        warehouse = str(location).strip() if location is not None else None
        date_iso = parse_date(date_raw)
        qty = parse_number(qty_raw)
        unit_cost = parse_number(price_raw)

        row_problems = []
        if not item_code:
            row_problems.append("Item Code kosong")
        elif not ITEM_CODE_RE.match(item_code):
            row_problems.append(f"Item Code mengandung karakter aneh: {item_code!r}")
        if not warehouse:
            row_problems.append("Warehouse/Location kosong")
        if date_iso is None:
            row_problems.append(f"Tanggal tidak bisa di-parse: {date_raw!r}")
        if qty is None:
            row_problems.append(f"Qty bukan angka: {qty_raw!r}")
        if unit_cost is None:
            row_problems.append(f"Price bukan angka: {price_raw!r}")

        if row_problems:
            warnings.append(f"  Baris {row_no}: " + "; ".join(row_problems))
            continue

        cleaned_rows.append([item_code, warehouse, date_iso, qty, unit_cost])
        unique_item_codes.add(item_code)

    if warnings:
        print("PERINGATAN - baris berikut dilewati karena data tidak valid:")
        print("\n".join(warnings))
        print()

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = "Opening Stock Import"
    out_ws.append(REQUIRED_COLUMNS)
    for row in cleaned_rows:
        out_ws.append(row)
    out_wb.save(out_path)

    print(f"Selesai. {len(cleaned_rows)} baris ditulis ke: {out_path}")
    print()
    print(f"Item Code unik yang perlu dicek kecocokannya dengan Item Master ({len(unique_item_codes)}):")
    for code in sorted(unique_item_codes):
        print(f"  - {code}")
    print()
    print("Catatan: template resmi Opening Stock TIDAK punya kolom UOM/Description/Item Batch --")
    print("field itu diambil dari Item Master saat import, bukan dari file. Satu baris = satu")
    print("dokumen Opening Stock (item+warehouse+tanggal+qty+cost), jadi file ini TIDAK perlu")
    print("dipecah per warehouse -- sistem menerima warehouse campuran dalam satu file.")


if __name__ == "__main__":
    main()
